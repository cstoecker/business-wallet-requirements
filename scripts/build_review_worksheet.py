#!/usr/bin/env python3
"""Build assets/downloads/ebw-review-worksheet.xlsx: the worksheet a human reviewer fills in.

Sheets: How to review; Requirements (P1 first, grouped by cluster, with dropdown decision columns);
Driver confirmations (criteria marked 'driver' on the architecture decision pages).
The filled file is read back by scripts/import_review_feedback.py.
"""
import os, re, yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
OUT = os.path.join(ROOT, "assets", "downloads"); os.makedirs(OUT, exist_ok=True)
def load(p): return yaml.safe_load(open(os.path.join(ROOT, p), encoding="utf-8")) or []
PET = "023852"; HEAD = PatternFill("solid", fgColor=PET); REV = PatternFill("solid", fgColor="DFFFF2")
reqs = {}
for f in os.listdir(os.path.join(ROOT, "_requirements")):
    t = open(os.path.join(ROOT, "_requirements", f), encoding="utf-8").read()
    m = re.match(r"---\n(.*?)\n---", t, re.S); r = yaml.safe_load(m.group(1)); reqs[r["req_id"]] = r
rv = {x["id"]: x for x in load("_data/review/requirement-review.yml")}
clusters = {c["code"]: c["name"] for c in load("_data/graph/clusters.yml")}
BASE = "https://spherity.github.io/business-wallet-requirements/requirements/"
wb = Workbook()
ws = wb.active; ws.title = "How to review"
how = [
 ("European Business Wallet requirements: review worksheet", ""),
 ("Purpose", "You confirm, change or reject draft requirements and tell us which architecture drivers are real. The result is read back by a script and applied to the catalogue. Analysis, not legal advice."),
 ("Order of work", "1. Sheet 'Requirements' is sorted by priority, then cluster. Start with P1 (mandatory and architecture-shaping). P2 and P3 only after P1, or by sample."),
 ("For every requirement ask", "a) Is the statement supported by the source and location (open the URL, read the clause, not only the extract)? b) Is it one obligation, unambiguous and testable? c) Right category, cluster and priority? d) Do you accept the review proposal in column 'Review proposal'?"),
 ("Decision (column M)", "agree = statement, category and priority are fine; agree with change = fill 'Your corrected statement' (must contain 'shall'); reject = not an obligation for a wallet or trust service design, or wrong; defer = needs more research; needs legal review = a lawyer must decide. Only 'agree' and 'agree with change' set the status to agreed."),
 ("Structural proposals", "merge, split, drop and recategorise change IDs. Do not edit IDs. Write your decision in column M and the target in column O; the maintainer carries it out."),
 ("Your name and date", "Fill columns Q and R on every row you decided. A requirement can only be 'agreed' with a named reviewer."),
 ("Sheet 'Driver confirmations'", "Criteria marked 'driver' on the architecture decision pages are assumptions until you confirm them. Say yes, no or modify, and give a weight from 1 (minor) to 5 (decisive)."),
 ("How to return it", "Save the file under your name, then either (1) commit it to the repository folder reviews/ on the branch you were given, (2) attach it to a GitHub discussion in the category 'Requirements review', or (3) send it to info@spherity.com. For single requirements you can also use the 'Give feedback' link on the requirement page (GitHub issue form)."),
 ("What happens next", "The maintainer runs scripts/import_review_feedback.py (dry run first), the requirement status, reviewer, statements and notes are updated, structural changes are listed for execution, and your feedback is stored in _data/review/feedback.yml with your name and the date."),
]
for a, b in how: ws.append([a, b])
ws["A1"].font = Font(name="Arial", bold=True, size=14, color=PET)
for row in ws.iter_rows(min_row=2):
    row[0].font = Font(name="Arial", bold=True, color=PET); row[1].font = Font(name="Arial"); row[1].alignment = Alignment(wrap_text=True, vertical="top"); row[0].alignment = Alignment(vertical="top")
ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 120
# requirements sheet
ws = wb.create_sheet("Requirements")
HDR = ["ID", "Priority", "Clusters", "Category", "Title", "Statement", "Source and location", "Source extract", "Review proposal", "Proposal note", "Proposed wording", "URL",
       "Decision", "Your corrected statement", "Target or change (category, priority, merge target)", "Comment", "Reviewer name", "Date"]
ws.append(HDR)
def key(i):
    x = rv.get(i, {}); c = (x.get("clusters") or ["K99"])[0]
    return (x.get("priority", "P3"), c, i)
for rid in sorted(reqs, key=key):
    r = reqs[rid]; x = rv.get(rid, {})
    prop = (x.get("verdict", "") + (" " + str(x["verdict_arg"]) if x.get("verdict_arg") else "")).strip()
    cs = x.get("corrected_statement")
    ws.append([rid, x.get("priority", ""), "; ".join(f"{c} {clusters.get(c, '')}" for c in x.get("clusters", [])), r["category"], r["title"], r["statement"],
               "; ".join(f"{s['id']} {s['location']}" for s in r["sources"]), r.get("full_quote") or r.get("quote", ""), prop, x.get("note") or x.get("priority_reason", ""),
               "\n".join(cs) if cs else "", BASE + rid.lower() + "/", "", "", "", "", "", ""])
widths = [14, 8, 34, 9, 30, 60, 36, 50, 18, 36, 50, 38, 18, 50, 28, 40, 18, 12]
for i, w in enumerate(widths, 1): ws.column_dimensions[chr(64 + i)].width = w
for c in ws[1]: c.fill = HEAD; c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[1][12:]: c.fill = PatternFill("solid", fgColor="14B98A")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in row[12:]: c.fill = REV
ws.freeze_panes = "F2"; ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type="list", formula1='"agree,agree with change,reject,defer,needs legal review"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"M2:M{ws.max_row}")
# drivers sheet
ws2 = wb.create_sheet("Driver confirmations")
ws2.append(["Decision", "Criterion", "Text", "Basis", "Confirm as requirement", "Weight 1-5", "Comment", "Reviewer name", "Date"])
dd = os.path.join(ROOT, "architecture", "decisions")
for f in sorted(os.listdir(dd)):
    if not f.startswith("dec-"): continue
    t = open(os.path.join(dd, f), encoding="utf-8").read(); dec = f[:-3].upper()
    for line in t.splitlines():
        m = re.match(r"\|\s*([A-Z]\d+)\s*\|\s*(.+?)\s*\|\s*(.+?)\s*\|$", line)
        if m and "driver" in m.group(3).lower():
            ws2.append([dec, m.group(1), m.group(2), m.group(3), "", "", "", "", ""])
for i, w in enumerate([10, 10, 70, 70, 22, 12, 40, 18, 12], 1): ws2.column_dimensions[chr(64 + i)].width = w
for c in ws2[1]: c.fill = HEAD; c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="top")
for row in ws2.iter_rows(min_row=2):
    for c in row: c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in row[4:]: c.fill = REV
dv2 = DataValidation(type="list", formula1='"yes,no,modify"', allow_blank=True); ws2.add_data_validation(dv2); dv2.add(f"E2:E{max(ws2.max_row, 2)}")
dv3 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); ws2.add_data_validation(dv3); dv3.add(f"F2:F{max(ws2.max_row, 2)}")
ws2.freeze_panes = "D2"
wb.save(os.path.join(OUT, "ebw-review-worksheet.xlsx"))
print("review worksheet written:", len(reqs), "requirements,", ws2.max_row - 1, "driver criteria")
