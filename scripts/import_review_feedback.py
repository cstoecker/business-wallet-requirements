#!/usr/bin/env python3
"""Read a filled review worksheet (scripts/build_review_worksheet.py) and apply it.

  scripts/import_review_feedback.py FILE.xlsx            dry run: prints what would change
  scripts/import_review_feedback.py FILE.xlsx --apply    updates _requirements/*.md and appends to _data/review/feedback.yml

Rules: a row without a Decision is ignored. 'agree' and 'agree with change' set status agreed and need a reviewer name.
'agree with change' needs a corrected statement with the word shall. 'reject' sets status obsolete, 'defer' and
'needs legal review' set status in-review. Merge, split, drop and recategorise proposals are never executed here: they are
listed for the maintainer. Driver confirmations go to _data/review/driver-confirmations.yml.
"""
import datetime, os, re, sys, yaml
from openpyxl import load_workbook
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
apply = "--apply" in sys.argv
files = [a for a in sys.argv[1:] if not a.startswith("--")]
rev = {x["id"]: x for x in yaml.safe_load(open(os.path.join(ROOT, "_data/review/requirement-review.yml"), encoding="utf-8"))}
def s(v): return "" if v is None else str(v).strip()
fb, drv, structural, problems, changed = [], [], [], [], 0
for path in files:
    wb = load_workbook(path, data_only=True)
    ws = wb["Requirements"]; hdr = [s(c.value) for c in ws[1]]; col = {h: i for i, h in enumerate(hdr)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid = s(row[col["ID"]]); dec = s(row[col["Decision"]]).lower()
        if not rid or not dec: continue
        who = s(row[col["Reviewer name"]]); date = s(row[col["Date"]])[:10] or datetime.date.today().isoformat()
        corr = s(row[col["Your corrected statement"]]); target = s(row[col["Target or change (category, priority, merge target)"]]); com = s(row[col["Comment"]])
        f = os.path.join(ROOT, "_requirements", rid + ".md")
        if not os.path.exists(f): problems.append(f"{rid}: unknown id"); continue
        if dec in ("agree", "agree with change") and not who: problems.append(f"{rid}: {dec} needs a reviewer name"); continue
        if dec == "agree with change" and not re.search(r"\bshall\b", corr): problems.append(f"{rid}: corrected statement must contain 'shall'"); continue
        if dec not in ("agree", "agree with change", "reject", "defer", "needs legal review"): problems.append(f"{rid}: unknown decision '{dec}'"); continue
        fb.append({"id": rid, "decision": dec, "corrected_statement": corr or None, "target": target or None, "comment": com or None, "reviewer": who or None, "date": date, "file": os.path.basename(path)})
        prop = s(row[col["Review proposal"]])
        if prop.split(" ")[0] in ("merge-with", "split", "drop", "recategorise") and dec in ("agree", "agree with change"):
            structural.append(f"{rid}: reviewer {dec}; proposal '{prop}'; target '{target}'")
        t = open(f, encoding="utf-8").read(); m = re.match(r"---\n(.*?)\n---\n?(.*)", t, re.S); fm = yaml.safe_load(m.group(1)); body = m.group(2)
        if dec == "agree": fm["status"] = "agreed"
        elif dec == "agree with change": fm["status"] = "agreed"; fm["statement"] = corr; fm["revision_note"] = "Statement changed by the reviewer."
        elif dec == "reject": fm["status"] = "obsolete"
        else: fm["status"] = "in-review"
        if who: fm["reviewer"] = who; fm["reviewed_on"] = date
        if com: fm["review_comment"] = com
        changed += 1
        print(("APPLY " if apply else "would ") + f"{rid}: {dec} -> status {fm['status']}" + (f" ({who})" if who else ""))
        if apply: open(f, "w", encoding="utf-8").write("---\n" + yaml.safe_dump(fm, sort_keys=False, allow_unicode=True, width=1000) + "---\n" + body)
    if "Driver confirmations" in wb.sheetnames:
        w2 = wb["Driver confirmations"]
        for r in w2.iter_rows(min_row=2, values_only=True):
            if s(r[4]):
                drv.append({"decision": s(r[0]), "criterion": s(r[1]), "confirmed": s(r[4]).lower(), "weight": r[5], "comment": s(r[6]) or None, "reviewer": s(r[7]) or None, "date": s(r[8])[:10] or datetime.date.today().isoformat()})
if apply:
    for name, items in (("feedback.yml", fb), ("driver-confirmations.yml", drv)):
        p = os.path.join(ROOT, "_data/review", name)
        old = yaml.safe_load(open(p, encoding="utf-8")) if os.path.exists(p) else []
        open(p, "w", encoding="utf-8").write("# Reviewer feedback imported by scripts/import_review_feedback.py\n" + yaml.safe_dump((old or []) + items, sort_keys=False, allow_unicode=True, width=1000))
print(f"{changed} requirements, {len(drv)} driver confirmations, {len(problems)} problems" + ("" if apply else " (dry run)"))
for x in problems: print(" problem:", x)
if structural:
    print("Structural proposals to execute by the maintainer:")
    for x in structural: print(" -", x)
