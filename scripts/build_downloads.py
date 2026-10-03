#!/usr/bin/env python3
"""Generate the downloads from the data: Excel workbook, CSV files and a JSON-LD graph in assets/downloads/.

Run before every Jekyll build (CI does this), so the Excel always matches the published requirements.
Sources of truth: _requirements/*.md (front matter), _data/graph/*.yml, _data/categories.yml.
"""
import csv, glob, json, os, re, sys, datetime
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
OUT = os.path.join(ROOT, "assets", "downloads"); os.makedirs(OUT, exist_ok=True)
def load(p): return yaml.safe_load(open(os.path.join(ROOT, p), encoding="utf-8")) or []
def reqs():
    rows = []
    for f in sorted(glob.glob(os.path.join(ROOT, "_requirements", "*.md"))):
        t = open(f, encoding="utf-8").read(); m = re.match(r"---\n(.*?)\n---", t, re.S)
        rows.append(yaml.safe_load(m.group(1)))
    return rows
R = reqs(); cats = load("_data/categories.yml"); src = load("_data/graph/sources.yml"); eco = load("_data/graph/ecosystems.yml")
cand = load("_data/graph/ecosystem-requirements.yml"); claims = load("_data/graph/claims.yml"); concepts = load("_data/graph/concepts.yml")
srcmap = {s["id"]: s for s in src}; today = datetime.date.today().isoformat()
PETROL, CYAN, OFF = "023852", "2BFEBB", "F4F2EE"
def j(v): return "; ".join(map(str, v)) if isinstance(v, list) else ("" if v is None else v)
def sheet(wb, title, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(); ws.title = title
    ws.append(header)
    for r in rows: ws.append(r)
    thin = Side(style="thin", color="C9C7C2")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color=OFF); c.fill = PatternFill("solid", fgColor=PETROL); c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=10); c.border = Border(bottom=thin)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    return ws
wb = Workbook()
ws = wb.active; ws.title = "Read me"
info = [("European Business Wallet Requirements: requirements catalogue", ""), ("Generated", today), ("Status", "Draft for review. Requirements are derived from the cited sources and have not yet been reviewed by a named person unless the Reviewer column says so. Based on the Commission proposal COM(2025) 838 where legal status is 'proposal'; it may change."),
        ("Analysis, not legal advice", "Legal content is analysis. Check article numbers against the Official Journal or the cited document."), ("Provenance", "L = law, S = standard or specification, D = derived by the project, A = assumption"),
        ("Sheets", "Requirements (derived, one per row); Candidates (requirement statements taken from ecosystems, not yet derived); Sources; Claims (business cases); Concepts; Categories"),
        ("Counts", f"{len(R)} requirements, {len(cand)} ecosystem candidates, {len(src)} sources, {len(claims)} claims, {len(concepts)} concepts"),
        ("Licence", "CC BY 4.0 (to be confirmed). Publisher: Spherity GmbH. Contact: info@spherity.com"), ("Web", "https://spherity.github.io/business-wallet-requirements/requirements/")]
for a, b in info: ws.append([a, b])
ws["A1"].font = Font(name="Arial", bold=True, size=14, color=PETROL)
for row in ws.iter_rows(min_row=2):
    row[0].font = Font(name="Arial", bold=True, color=PETROL); row[1].font = Font(name="Arial"); row[1].alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 110
catname = {c["code"]: c["name"] for c in cats}
REV = {x["id"]: x for x in (load("_data/review/requirement-review.yml") if os.path.exists(os.path.join(ROOT, "_data/review/requirement-review.yml")) else [])}
CLU = {c["code"]: c["name"] for c in load("_data/graph/clusters.yml")}
rr = []
for r in sorted(R, key=lambda x: x["req_id"]):
    rr.append([r["req_id"], r["category"], catname.get(r["category"], ""), r["title"], r["statement"], r["rationale"],
               j([f"{s['id']} ({s['location']})" for s in r["sources"]]), r["provenance"], r["legal_status"], r["plane"], j(r["perspectives"]), j(r["actors"]),
               r["verification_method"], r["status"], r["reviewer"], j(r.get("concepts")), r.get("created"), r.get("last_verified"),
               "https://spherity.github.io/business-wallet-requirements/requirements/" + r["req_id"].lower() + "/",
               r.get("quote", ""), REV.get(r["req_id"], {}).get("priority", ""), j([f"{c} {CLU.get(c, '')}" for c in REV.get(r["req_id"], {}).get("clusters", [])]), REV.get(r["req_id"], {}).get("applicability", ""),
               (REV.get(r["req_id"], {}).get("verdict", "") + (" " + str(REV[r["req_id"]]["verdict_arg"]) if REV.get(r["req_id"], {}).get("verdict_arg") else "")).strip(), REV.get(r["req_id"], {}).get("note", "")])
HDR = ["ID", "Category", "Category name", "Title", "Requirement", "Rationale", "Sources", "Provenance", "Legal status", "Plane", "Perspectives", "Actors", "Verification method", "Status", "Reviewer", "Concepts", "Created", "Last verified", "URL", "Source extract", "Priority", "Clusters", "Applicability", "Review proposal", "Review note"]
sheet(wb, "Requirements", HDR, rr, [14, 9, 22, 30, 60, 45, 45, 11, 12, 12, 16, 28, 14, 10, 10, 28, 12, 12, 40, 50, 9, 36, 16, 16, 40])
sheet(wb, "Candidates", ["ID", "Ecosystem", "Category", "Statement", "Reference", "Source", "Verification", "Status"],
      [[c["id"], c["ecosystem"], c["category"], c["statement"], c["ref"], c["source"], c["verification"], c["status"]] for c in cand], [14, 18, 9, 70, 40, 22, 12, 11])
sheet(wb, "Sources", ["ID", "Title", "Issuer", "Class", "Version", "Date", "URL", "Retrieved", "Verification", "Note"],
      [[s["id"], s["title"], s["issuer"], s["class"], s.get("version", ""), s.get("date", ""), s.get("url", ""), s.get("retrieved", ""), s.get("verification", ""), s.get("note", "")] for s in src], [24, 60, 30, 18, 18, 12, 50, 12, 12, 40])
sheet(wb, "Claims", ["ID", "Business case", "Statement", "Figure", "Unit", "Basis", "Geography", "Type", "Source", "Location", "Verification", "Caveat"],
      [[c["id"], c["case"], c["statement"], c.get("figure", ""), c.get("unit", ""), c.get("basis", ""), c.get("geography", ""), c.get("type", ""), c["source"], c.get("location", ""), c["verification"], c.get("caveat", "")] for c in claims], [13, 14, 70, 9, 14, 20, 14, 12, 24, 28, 12, 50])
sheet(wb, "Concepts", ["ID", "Title", "Group", "Status", "Summary", "Related", "Categories"],
      [[c["id"], c["title"], c["group"], c["status"], c["summary"], j(c.get("related")), j(c.get("requirement_categories"))] for c in concepts], [24, 40, 20, 10, 70, 40, 14])
sheet(wb, "Categories", ["Code", "Name", "Scope", "Plane", "Requirements", "Candidates"],
      [[c["code"], c["name"], c["scope"], c["plane"], sum(1 for r in R if r["category"] == c["code"]), sum(1 for x in cand if x["category"] == c["code"])] for c in cats], [8, 32, 80, 24, 14, 12])
wb.save(os.path.join(OUT, "ebw-requirements.xlsx"))
def csvw(name, header, rows):
    with open(os.path.join(OUT, name), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
csvw("requirements.csv", HDR, rr)
csvw("sources.csv", ["id", "title", "issuer", "class", "version", "date", "url", "retrieved", "verification"], [[s["id"], s["title"], s["issuer"], s["class"], s.get("version", ""), s.get("date", ""), s.get("url", ""), s.get("retrieved", ""), s.get("verification", "")] for s in src])
csvw("claims.csv", ["id", "case", "statement", "figure", "unit", "basis", "geography", "type", "source", "location", "verification", "caveat"], [[c["id"], c["case"], c["statement"], c.get("figure", ""), c.get("unit", ""), c.get("basis", ""), c.get("geography", ""), c.get("type", ""), c["source"], c.get("location", ""), c["verification"], c.get("caveat", "")] for c in claims])
BASE = "https://spherity.github.io/business-wallet-requirements/"
g = {"@context": {"@vocab": "https://schema.org/", "ebw": BASE + "ns#"}, "@graph": []}
for r in R:
    g["@graph"].append({"@id": BASE + "requirements/" + r["req_id"].lower() + "/", "@type": ["ebw:Requirement", "DefinedTerm"], "identifier": r["req_id"], "name": r["title"], "description": r["statement"], "ebw:category": r["category"],
                        "ebw:provenance": r["provenance"], "ebw:status": r["status"], "ebw:derivedFrom": [{"@id": srcmap[s["id"]]["url"]} for s in r["sources"] if s["id"] in srcmap]})
for s in src: g["@graph"].append({"@id": s.get("url") or BASE + "sources/" + s["id"], "@type": "ebw:Source", "identifier": s["id"], "name": s["title"], "publisher": s["issuer"], "version": s.get("version", "")})
for c in concepts: g["@graph"].append({"@id": BASE + "concepts/" + c["slug"] + "/", "@type": "DefinedTerm", "identifier": c["id"], "name": c["title"], "description": c["summary"]})
json.dump(g, open(os.path.join(OUT, "graph.jsonld"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
print(f"downloads written: {len(R)} requirements, {len(cand)} candidates, {len(src)} sources, {len(claims)} claims")
